#!/usr/bin/env python3
"""Adds Chili's, Burger King, Raising Cane's and Chipotle to restaurants.json from a spreadsheet.

These chains don't offer a file the weekly workflow can download (interactive menus, bot
protection, or a private ordering API), so their numbers were collected by hand from each
chain's official nutrition page into data/*.xlsx (one sheet per chain, plus a Notes sheet).
To refresh: drop in a new spreadsheet with the same columns and run
    pip install openpyxl
    python3 scripts/import-sheet.py data/<file>.xlsx
The weekly restaurants workflow keeps chains it doesn't fetch itself, so this data stays put.
"""
import json, re, sys
from datetime import date
from pathlib import Path
import openpyxl

OUT = Path(__file__).resolve().parent.parent / 'restaurants.json'
CHAINS = {
    'Chilis': ('chilis', "Chili's", ['chilis', "chili's"]),
    'Burger King': ('burgerking', 'Burger King', ['burger king', 'bk']),
    'Raising Canes': ('raisingcanes', "Raising Cane's", ['raising canes', 'canes', "cane's"]),
    'Chipotle': ('chipotle', 'Chipotle', ['chipotle']),
}
# Not a meal for one person (and the app is used by teenagers).
SKIP_CATEGORIES = {'Alcohol', 'Party Platters'}

def num(v):
    if isinstance(v, (int, float)): return v
    s = str(v or '').strip()
    if s.startswith('<'): return 0
    try: return float(s)
    except ValueError: return None

def clean(s):
    return re.sub(r'\s+', ' ', str(s).replace('®', '').replace('™', '').replace('’', "'")).strip()

def rows(ws):
    out = []
    for r in ws.iter_rows(min_row=2, values_only=True):
        if not r or not r[2]: continue
        _, cat, name, serving, cal, f, c, p = r[:8]
        if cat in SKIP_CATEGORIES: continue
        item = dict(name=clean(name), category=cat, serving=clean(serving or ''),
                    cal=num(cal), f=num(f), c=num(c), p=num(p))
        if item['cal'] is None: continue
        if re.search(r'\((Louisville|Maryland)', item['name']): continue  # regional versions of the kids' combo
        out.append(item)
    return out

def chipotle_builds(items):
    """Common orders built from Chipotle's per-ingredient numbers (adults' portions)."""
    ing = {i['name']: i for i in items if i['category'] == 'Entree Ingredients'}
    def build(name, parts, note):
        got = [ing.get(p) for p in parts]
        if not all(got): return None
        t = {k: round(sum(g[k] or 0 for g in got), 1) for k in ('cal', 'f', 'c', 'p')}
        return dict(name=name, serving=note, **t)
    base = ['Cilantro-Lime White Rice', 'Black Beans', 'Fresh Tomato Salsa']
    out = []
    for prot in ['Chicken', 'Steak', 'Barbacoa', 'Carnitas', 'Sofritas']:
        for b in [
            build(f'{prot} Bowl', [*base, prot], 'white rice, black beans, mild salsa'),
            build(f'{prot} Bowl with Cheese & Sour Cream', [*base, prot, 'Cheese', 'Sour Cream'], 'white rice, black beans, mild salsa, cheese, sour cream'),
            build(f'{prot} Bowl, Double {prot}', [*base, prot, prot], f'white rice, black beans, mild salsa, double {prot.lower()}'),
            build(f'{prot} Burrito', ['Flour Tortilla (burrito)', *base, prot, 'Cheese'], 'tortilla, white rice, black beans, mild salsa, cheese'),
            build(f'{prot} Salad', ['Supergreens Salad Mix', 'Black Beans', 'Fresh Tomato Salsa', prot], 'supergreens, black beans, mild salsa (no dressing)'),
            build(f'{prot} Tacos (3, soft)', ['Flour Tortilla (taco)'] * 3 + [prot, 'Fresh Tomato Salsa', 'Cheese', 'Romaine Lettuce (tacos)'], '3 flour tortillas, meat, mild salsa, cheese, lettuce'),
        ]:
            if b: out.append(b)
    return out

def main():
    wb = openpyxl.load_workbook(sys.argv[1], data_only=True)
    notes = {}
    if 'Notes' in wb.sheetnames:
        for r in wb['Notes'].iter_rows(min_row=2, values_only=True):
            if r and r[0]: notes[clean(r[0])] = clean(r[1] or '')
    data = json.loads(OUT.read_text())
    for sheet, (cid, name, aliases) in CHAINS.items():
        if sheet not in wb.sheetnames: continue
        items = rows(wb[sheet])
        extra = chipotle_builds(items) if cid == 'chipotle' else []
        if cid == 'chipotle':
            for i in items:
                if i['category'] == 'Kids Menu': i['name'] = f"Kids {i['name']}"
                elif i['category'] == 'Entree Ingredients' and i['serving']:
                    # "Chicken" -> "Chicken (4 oz)"; "Flour Tortilla (burrito)" -> "Flour Tortilla (burrito, 1 ea)"
                    n = i['name']
                    i['name'] = f"{n[:-1]}, {i['serving']})" if n.endswith(')') else f"{n} ({i['serving']})"
        # Same name twice (drink sizes): tell them apart by serving, then by category.
        counts = {}
        for i in items: counts[i['name'].lower()] = counts.get(i['name'].lower(), 0) + 1
        for i in items:
            if counts[i['name'].lower()] > 1:
                i['name'] = f"{i['name']} ({i['serving'] or i['category']})"
        seen, final = set(), []
        for i in items + extra:
            k = i['name'].lower()
            if k in seen: continue
            seen.add(k)
            final.append({k2: i[k2] for k2 in ('name', 'serving', 'cal', 'f', 'c', 'p') if k2 in i})
        src = notes.get(name) or notes.get(sheet) or 'official nutrition page'
        data['chains'][cid] = {'name': name, 'aliases': aliases, 'source': f'{src} (data/{Path(sys.argv[1]).name})',
                               'updated': date.today().isoformat(), 'items': final}
        data.setdefault('status', {})[cid] = f'ok: {len(final)} items (manual import)'
        print(f'{name}: {len(final)} items')
    OUT.write_text(json.dumps(data, indent=1, ensure_ascii=False) + '\n')

if __name__ == '__main__':
    main()
